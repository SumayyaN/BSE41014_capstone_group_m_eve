"""Week 7 live runner for the E2E scenarios in Group_M_Week7_Evaluation_Sheet_FINAL.xlsx.
(Separate from the Week 3 RAG evaluation script run_eval.py.)

Run from the repo's src/ folder, in the environment that starts app.py:

    python run_scenarios.py Group_M_Week7_Evaluation_Sheet_FINAL.xlsx --repeat 1 --out results_live.csv
    python run_scenarios.py Group_M_Week7_Evaluation_Sheet_FINAL.xlsx --only B01,B02,N01,A01 --repeat 1

For each E2E scenario it creates a fresh SessionState and a temporary
drafts.json (seeded from the preconditions column), sets case_agreed before
the turn named in sidebar_agree_before_turn (same effect as the sidebar
button), and runs each student turn through run_tool_turn.

It writes
  * a CSV row per turn (reply, stop condition, tools, latency, drafts added), and
  * one JSON trace per run in evidence/traces/week7/, in the same shape as the
    Week 5 traces, plus model id, prompt version, per-call latency and
    retrieval scores.

It does NOT judge the reply text. Use must_include / must_not_include in the
workbook to fill groundedness, tool_selection_ok, safety_violation and pass_fail.

Quota: the free tier allowed about 20 model requests a day in Week 4. A scenario
that stops in Sense (S4/S5/S6) uses none. Use --only to run in batches, and
--pause to space calls out. Replies that start with "Error:" and mention 429 or
503 are treated as infrastructure errors, retried after --backoff seconds, and
flagged in the CSV so they are not counted as product failures.
"""
import argparse, csv, json, pathlib, tempfile, time
from datetime import datetime, timezone
from openpyxl import load_workbook

from agent.state import SessionState
from tools import case_store
import tools.orchestrate as orch

SEEDS = {
    "none": [],
    "SEED_A": [{"case_id": "CASE-2026-000123", "session_id": "seed-1", "student_identifier": "26/U/10003",
                "issue_summary": "Cannot activate webmail after following the SSO steps", "category": "it_webmail_wifi",
                "queue_label": "ict_helpdesk", "created_at": "2026-10-08T09:15:00+00:00"}],
    "SEED_B": [{"case_id": "CASE-2026-000200", "session_id": "seed-2", "student_identifier": "26/U/10009",
                "issue_summary": "Cannot activate webmail", "category": "it_webmail_wifi",
                "queue_label": "ict_helpdesk", "created_at": "2026-10-01T08:00:00+00:00"},
               {"case_id": "CASE-2026-000201", "session_id": "seed-3", "student_identifier": "26/U/10009",
                "issue_summary": "Registration centre unclear for my college", "category": "registration",
                "queue_label": "registration_office", "created_at": "2026-10-05T08:00:00+00:00"}],
}
PROMPT_VERSION = "v1.3"
MODEL_CALLS: list[dict] = []   # filled by the timing wrapper, cleared per turn
TOOL_TIMINGS: list[dict] = []


# --- timing wrappers (no change to the app's own code) -----------------------
def _install_timers():
    # Model calls: record latency for successful calls too.
    inner_models = orch.client.models
    orig = inner_models.generate_content

    def timed_generate(**kwargs):
        start = time.perf_counter()
        try:
            out = orig(**kwargs)
            MODEL_CALLS.append({"latency_ms": round((time.perf_counter() - start) * 1000)})
            return out
        except Exception as exc:
            MODEL_CALLS.append({"latency_ms": round((time.perf_counter() - start) * 1000), "error": str(exc)[:200]})
            raise

    class _Wrap:
        def __init__(self, models):
            self._m = models
        generate_content = staticmethod(timed_generate)
        def __getattr__(self, n):
            return getattr(self._m, n)

    class _Cli:
        def __init__(self, cli):
            self._c = cli
            self.models = _Wrap(cli.models)
        def __getattr__(self, n):
            return getattr(self._c, n)

    orch.client = _Cli(orch.client)

    orig_run = orch._run_call

    def timed_run(name, arguments, state, question):
        start = time.perf_counter()
        result = orig_run(name, arguments, state, question)
        TOOL_TIMINGS.append({"tool": name, "latency_ms": round((time.perf_counter() - start) * 1000),
                             "status": result.get("status")})
        return result

    orch._run_call = timed_run


# --- helpers -------------------------------------------------------------------
def _short_result(result: dict) -> dict:
    status = result.get("status")
    brief = {"status": status}
    for key in ("message", "validation_errors", "retrieved_count", "case_id", "session_id", "routed_queue", "created_at"):
        if result.get(key) is not None:
            brief[key] = result[key]
    if status == "ok":
        brief["sources"] = [{"source_title": c.get("source_title"), "source_url": c.get("source_url"),
                             "similarity_score": c.get("similarity_score")} for c in result.get("chunks") or []]
    return brief


def load_scenarios(path):
    """Read the 'Runner input' sheet of the final evaluation workbook."""
    ws = load_workbook(path, data_only=True)["Runner input"]
    head = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2, values_only=True):
        rec = dict(zip(head, row))
        if rec.get("id") and rec.get("student_turns"):
            yield rec


def _is_infra_error(text: str) -> bool:
    if not text:
        return False
    if text.startswith("The assistant is temporarily busy."):
        return True
    return text.startswith("Error:") and any(
        k in text for k in ("429", "503", "UNAVAILABLE", "RESOURCE_EXHAUSTED")
    )


def run_once(rec, pause):
    store = pathlib.Path(tempfile.mkdtemp()) / "drafts.json"
    seed = SEEDS.get(rec["preconditions"] or "none", [])
    store.write_text(json.dumps(seed))
    case_store.STORE_PATH = store
    before = len(seed)

    state = SessionState()
    agree_at = rec["sidebar_agree_before_turn"]
    turns = [t.strip() for t in rec["student_turns"].split("||")]
    out, infra = [], False
    for n, msg in enumerate(turns, 1):
        events = []
        if agree_at and n == int(agree_at):
            state.case_agreed = True
            events.append("Student used the agreement control. case_agreed is now true.")
        MODEL_CALLS.clear(); TOOL_TIMINGS.clear()
        start = time.perf_counter()
        r = orch.run_tool_turn(state, msg)
        total_ms = round((time.perf_counter() - start) * 1000)
        infra = infra or _is_infra_error(r["text"])
        out.append({
            "application_events": events, "turn": n, "student_message": msg, "reply": r["text"],
            "stop_condition": r.get("stop_condition"), "iteration_count": r.get("iteration_count"),
            "latency_ms": total_ms, "model_calls": list(MODEL_CALLS), "tool_timings": list(TOOL_TIMINGS),
            "tools": [{"tool": s["tool"], "arguments": s["arguments"], "result": _short_result(s["result"])}
                      for s in r.get("trace") or []],
        })
        if pause:
            time.sleep(pause)
    after = json.loads(store.read_text())
    return state, out, len(after) - before, [x["case_id"] for x in after[before:]], infra


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook")
    ap.add_argument("--repeat", type=int, default=1)
    ap.add_argument("--out", default="results_live.csv")
    ap.add_argument("--only", default="", help="comma-separated scenario ids")
    ap.add_argument("--trace-dir", default="../evidence/traces/week7")
    ap.add_argument("--pause", type=float, default=0.0, help="seconds to wait after each turn")
    ap.add_argument("--retries", type=int, default=1, help="retries after an infrastructure error")
    ap.add_argument("--backoff", type=float, default=60.0)
    a = ap.parse_args()
    only = {x.strip() for x in a.only.split(",") if x.strip()}
    trace_dir = pathlib.Path(a.trace_dir); trace_dir.mkdir(parents=True, exist_ok=True)

    _install_timers()

    with open(a.out, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["id", "run", "turn", "student_message", "reply", "stop_actual", "iterations", "tools_called",
                    "latency_s", "model_calls", "drafts_added", "new_case_ids", "infra_error", "trace_file"])
        for rec in load_scenarios(a.workbook):
            if only and rec["id"] not in only:
                continue
            for k in range(1, a.repeat + 1):
                attempt, result = 0, None
                while True:
                    try:
                        result = run_once(rec, a.pause)
                    except Exception as exc:
                        w.writerow([rec["id"], k, "", "", f"RUNNER ERROR: {exc}", "", "", "", "", "", "", "", "yes", ""])
                        result = None
                        break
                    if not result[4] or attempt >= a.retries:
                        break
                    attempt += 1
                    print(f"{rec['id']} run {k}: infrastructure error, waiting {a.backoff}s (retry {attempt})")
                    time.sleep(a.backoff)
                if result is None:
                    continue
                state, turns, added, ids, infra = result
                fname = f"{rec['id']}_run{k}.json"
                (trace_dir / fname).write_text(json.dumps({
                    "id": rec["id"], "scenario_type": rec.get("scenario_type"),
                    "run": k, "recorded_at": datetime.now(timezone.utc).isoformat(),
                    "model": orch.MODEL, "prompt_version": PROMPT_VERSION,
                    "session_id": state.session_id, "preconditions": rec["preconditions"],
                    "drafts_added": added, "new_case_ids": ids, "infra_error": infra, "turns": turns,
                }, indent=2), encoding="utf-8")
                for t in turns:
                    last = t["turn"] == len(turns)
                    w.writerow([rec["id"], k, t["turn"], t["student_message"], t["reply"], t["stop_condition"],
                                t["iteration_count"], "; ".join(f"{x['tool']}:{x['result'].get('status')}" for x in t["tools"]),
                                round(t["latency_ms"] / 1000, 2), len(t["model_calls"]),
                                added if last else "", ", ".join(ids) if last else "", "yes" if infra and last else "", fname])
                fh.flush()
                last_t = turns[-1]
                print(f"{rec['id']} run {k}: stop={last_t['stop_condition']} tools={[x['tool'] for x in last_t['tools']]} "
                      f"latency={last_t['latency_ms']/1000:.1f}s drafts+{added}{' INFRA' if infra else ''}")


if __name__ == "__main__":
    main()
